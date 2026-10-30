"""Material review must check actual pages without inventing unavailable checks."""

from __future__ import annotations

import asyncio
import hashlib
import io
import json
import re
import zipfile
from pathlib import Path
from types import SimpleNamespace

import fitz
import pytest
from pptx import Presentation
from pptx.util import Inches

from app.agent.tools import material_assets
from app.agent.tools import material_review as review


@pytest.fixture
def managed(tmp_path, monkeypatch):
    config = SimpleNamespace(
        resolve_path=lambda raw: (tmp_path / raw).resolve(), api_prefix="/api/v1",
    )
    monkeypatch.setattr(review, "settings", config)
    monkeypatch.setattr(material_assets, "settings", config)
    monkeypatch.setattr(review, "_find_libreoffice", lambda: None)
    monkeypatch.setattr(review, "_resolve_review_model", lambda: SimpleNamespace(
        source="runtime", configured=True, model="test-vision",
    ))
    monkeypatch.setattr(review, "_supports_vision", lambda config: False)
    exports = config.resolve_path("./data/exports")
    exports.mkdir(parents=True)
    return exports


def make_pdf(path: Path, count=1, text="A complete training overview with examples and explanations."):
    document = fitz.open()
    for number in range(count):
        page = document.new_page()
        page.insert_text((55, 65), f"Page {number + 1}: {text}")
    document.save(path)
    document.close()
    return path


def make_pptx(path: Path, count=2):
    presentation = Presentation()
    for number in range(count):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1))
        shape.text = f"Lesson {number + 1}: a complete training overview"
        slide.notes_slide.notes_text_frame.text = f"Presenter notes for lesson {number + 1}"
    table = presentation.slides[0].shapes.add_table(
        2, 2, Inches(1), Inches(3), Inches(6), Inches(1),
    ).table
    table.cell(0, 0).text = "Thread A"
    table.cell(0, 1).text = "Thread B"
    table.cell(1, 0).text = "Read 0"
    table.cell(1, 1).text = "Read 0"
    presentation.save(path)
    return path


def run(**kwargs):
    return asyncio.run(review.review_materials(**kwargs))


def test_structural_pdf_review_reads_every_page_and_preserves_hash(managed):
    document = make_pdf(managed / "handbook.pdf", 3)
    result = run(files=[str(document)], visual=False)
    assert result["ok"] and result["review_complete"]
    assert result["review_scope"] == "structural_only"
    assert result["visual_review_complete"] is False
    file = result["files"][0]
    assert file["sha256"] == hashlib.sha256(document.read_bytes()).hexdigest()
    assert file["current_sha256"] == file["sha256"]
    assert file["page_count"] == 3
    assert [page["page"] for page in file["pages"]] == [1, 2, 3]
    assert all(page["preview_url"].startswith("/api/v1/files/exports/review_")
               and page["width"] <= review.MAX_RENDER_EDGE
               and page["height"] <= review.MAX_RENDER_EDGE for page in file["pages"])
    assert file["code_execution"]["status"] == "not_run"
    report = json.loads(Path(result["path"]).read_text(encoding="utf-8"))
    assert report["files"][0]["sha256"] == file["sha256"]
    assert report["review_scope"] == "structural_only"
    assert "未执行视觉审查" in report["message"]
    assert result["download_url"].endswith(Path(result["path"]).name)
    assert "preview_path" not in report["files"][0]["pages"][0]


def test_nonvisual_model_cannot_pass_visual_review(managed):
    path = make_pdf(managed / "lesson.pdf")
    result = run(files=[str(path)], visual=True)
    assert result["ok"]
    assert result["status"] == "partial"
    assert result["review_complete"] is False
    assert result["files"][0]["render"]["status"] == "rendered"
    assert result["files"][0]["visual"]["status"] == "unavailable"


def test_pptx_without_renderer_keeps_all_text_notes_and_native_tables(managed):
    path = make_pptx(managed / "lecture.pptx")
    result = run(files=[str(path)], expected_slide_count=2)
    file = result["files"][0]
    assert file["page_count"] == 2
    assert file["pages"][0]["table_count"] == 1
    assert "Presenter notes" in file["pages"][0]["notes_preview"]
    assert "Thread A" in file["pages"][0]["text_preview"]
    assert file["render"]["status"] == "unavailable"
    assert "LibreOffice" in file["render"]["reason"]
    assert result["review_complete"] is False


def test_exact_total_slide_count_includes_cover(managed):
    path = make_pptx(managed / "lecture.pptx", 2)
    result = run(files=[str(path)], expected_slide_count=8, visual=False)
    assert result["status"] == "needs_revision"
    assert not result["review_complete"]
    assert any("8" in issue["message"] and "2" in issue["message"] for issue in result["issues"])


def test_required_content_and_code_absence_are_actionable(managed):
    path = make_pdf(managed / "lesson.pdf")
    result = run(files=[str(path)], required_sections=["修正代码", "答案"], visual=False)
    assert result["status"] == "needs_revision"
    assert result["has_code"] is False
    assert any("修正代码" in issue["message"] for issue in result["issues"])
    assert any("未检测到非空代码" in issue["message"] for issue in result["issues"])


def test_rejects_files_outside_managed_directories(managed, tmp_path):
    outside = make_pdf(tmp_path / "private.pdf")
    result = run(files=[str(outside)])
    assert result["ok"] is False
    assert not list(managed.glob("review_*"))


def test_rejects_traversal_and_remote_urls(managed):
    for value in ("/api/v1/files/exports/../private.pdf", "https://example.com/handbook.pdf"):
        assert run(files=[value])["ok"] is False


def test_total_page_budget_checked_before_any_render(managed):
    one = make_pdf(managed / "part1.pdf", 31)
    two = make_pdf(managed / "part2.pdf", 30)
    result = run(files=[str(one), str(two)], visual=False)
    assert result["ok"] is False
    assert "60" in result["error"]
    assert not list(managed.glob("review_*"))


def test_visual_review_covers_every_page_in_bounded_batches(managed, monkeypatch):
    path = make_pdf(managed / "lesson.pdf", 5)
    monkeypatch.setattr(review, "_supports_vision", lambda config: True)
    batches = []

    async def model(messages, **kwargs):
        content = messages[-1]["content"]
        pages = [int(re.fullmatch(r"第 (\d+) 页", part["text"])[1])
                 for part in content if part["type"] == "text" and part["text"].startswith("第 ")]
        batches.append(pages)
        assert kwargs["config"].model == "test-vision"
        assert len([part for part in content if part["type"] == "image_url"]) <= 4
        return SimpleNamespace(content=json.dumps({
            "pages_reviewed": pages, "issues": [], "summary": "All supplied pages are readable.",
        }), finish_reason="stop")

    monkeypatch.setattr("app.llm.client.chat_complete", model)
    result = run(files=[str(path)])
    assert result["review_complete"] and result["visual_review_complete"]
    assert batches == [[1, 2, 3, 4], [5]]
    assert result["files"][0]["visual"]["checked_pages"] == [1, 2, 3, 4, 5]


def test_model_cannot_omit_a_page_and_still_pass(managed, monkeypatch):
    path = make_pdf(managed / "lesson.pdf", 2)
    monkeypatch.setattr(review, "_supports_vision", lambda config: True)

    async def model(*args, **kwargs):
        return SimpleNamespace(content=json.dumps({
            "pages_reviewed": [1], "issues": [], "summary": "Looks good.",
        }), finish_reason="stop")

    monkeypatch.setattr("app.llm.client.chat_complete", model)
    result = run(files=[str(path)])
    assert not result["review_complete"]
    assert result["files"][0]["visual"]["status"] == "unavailable"


def test_model_failure_is_an_unavailable_check_not_an_approval(managed, monkeypatch):
    path = make_pdf(managed / "lesson.pdf")
    monkeypatch.setattr(review, "_supports_vision", lambda config: True)

    async def model(*args, **kwargs):
        raise RuntimeError("A provider request failed")

    monkeypatch.setattr("app.llm.client.chat_complete", model)
    result = run(files=[str(path)])
    assert result["status"] == "partial"
    assert result["review_complete"] is False
    assert result["files"][0]["visual"]["unreviewed_pages"] == [1]
    saved = Path(result["path"]).read_text(encoding="utf-8")
    assert "data:image/png;base64" not in saved
    assert "api_key" not in saved


@pytest.mark.parametrize("prefix,suffix", [
    ("", ""), ("```json\n", "\n```"),
    ("以下是逐页审查结果：\n```json\n", "\n```"),
    ("审查结果如下：\n", ""), ("```json ", "```"),
])
def test_visual_json_accepts_complete_provider_fences_and_preambles(managed, monkeypatch, prefix, suffix):
    path = make_pdf(managed / "lesson.pdf")
    monkeypatch.setattr(review, "_supports_vision", lambda config: True)

    async def model(*args, **kwargs):
        return SimpleNamespace(content=prefix + json.dumps({
            "pages_reviewed": [1], "issues": [], "summary": "All supplied pages are readable.",
        }) + suffix, finish_reason="stop")

    monkeypatch.setattr("app.llm.client.chat_complete", model)
    result = run(files=[str(path)])
    assert result["visual_review_complete"] and result["review_complete"]
    assert result["files"][0]["visual"]["checked_pages"] == [1]


@pytest.mark.parametrize("finish_reason", ["length", "max_tokens", "interrupted", "content_filter"])
def test_valid_looking_json_with_incomplete_finish_reason_never_passes(managed, monkeypatch, finish_reason):
    path = make_pdf(managed / "lesson.pdf")
    monkeypatch.setattr(review, "_supports_vision", lambda config: True)

    async def model(*args, **kwargs):
        return SimpleNamespace(content=json.dumps({
            "pages_reviewed": [1], "issues": [], "summary": "Readable.",
        }), finish_reason=finish_reason)

    monkeypatch.setattr("app.llm.client.chat_complete", model)
    result = run(files=[str(path)])
    assert result["review_complete"] is False
    assert result["files"][0]["visual"]["checked_pages"] == []


@pytest.mark.parametrize("corrupt", [
    lambda valid: valid[:-1],
    lambda valid: valid + '\n{"more":',
    lambda valid: valid + '\n{"pages_reviewed":[1],"issues":[],"summary":"Second object"}',
    lambda valid: '{"unfinished_outer":' + valid,
])
def test_visual_parser_does_not_salvage_partial_or_ambiguous_objects(managed, monkeypatch, corrupt):
    path = make_pdf(managed / "lesson.pdf")
    monkeypatch.setattr(review, "_supports_vision", lambda config: True)

    async def model(*args, **kwargs):
        valid = json.dumps({"pages_reviewed": [1], "issues": [], "summary": "Readable."})
        return SimpleNamespace(content=corrupt(valid), finish_reason="stop")

    monkeypatch.setattr("app.llm.client.chat_complete", model)
    result = run(files=[str(path)])
    assert result["status"] == "partial"
    assert result["visual_review_complete"] is False
    assert result["files"][0]["visual"]["unreviewed_pages"] == [1]


def test_a_visual_error_requires_revision(managed, monkeypatch):
    path = make_pdf(managed / "lesson.pdf")
    monkeypatch.setattr(review, "_supports_vision", lambda config: True)

    async def model(*args, **kwargs):
        return SimpleNamespace(content=json.dumps({
            "pages_reviewed": [1], "issues": [
                {"page": 1, "severity": "error", "message": "Code overlaps the footer; split this page."}
            ], "summary": "Layout needs revision.",
        }), finish_reason="stop")

    monkeypatch.setattr("app.llm.client.chat_complete", model)
    result = run(files=[str(path)])
    assert result["status"] == "needs_revision"
    assert not result["review_complete"]
    assert result["issues"][0]["page"] == 1


def test_hash_detects_a_file_changed_after_review_started(managed, monkeypatch):
    path = make_pdf(managed / "lesson.pdf")
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    monkeypatch.setattr(review, "_supports_vision", lambda config: True)

    async def model(*args, **kwargs):
        path.write_bytes(path.read_bytes() + b"\n%updated\n")
        return SimpleNamespace(content=json.dumps({
            "pages_reviewed": [1], "issues": [], "summary": "Readable.",
        }), finish_reason="stop")

    monkeypatch.setattr("app.llm.client.chat_complete", model)
    result = run(files=[str(path)])
    file = result["files"][0]
    assert file["sha256"] == before
    assert file["current_sha256"] != before
    assert result["status"] == "needs_revision"
    assert any("文件已更改" in issue["message"] for issue in result["issues"])


def test_libreoffice_runs_with_isolated_profile_and_no_shell(managed, monkeypatch):
    path = make_pptx(managed / "lecture.pptx")
    monkeypatch.setattr(review, "_find_libreoffice", lambda: "soffice-test")
    invoked = []

    def convert(args, **kwargs):
        invoked.append(args)
        assert kwargs["shell"] is False
        assert kwargs["timeout"] == review.RENDER_TIMEOUT_SECONDS
        assert "--headless" in args and "--norestore" in args
        assert any(value.startswith("-env:UserInstallation=file:") for value in args)
        work = Path(kwargs["cwd"])
        assert (work / "profile/user/registrymodifications.xcu").is_file()
        assert (work / "source.pptx").read_bytes() == path.read_bytes()
        make_pdf(work / "source.pdf", 2)
        return SimpleNamespace(returncode=0)

    monkeypatch.setattr(review.subprocess, "run", convert)
    result = run(files=[str(path)], expected_slide_count=2, visual=False)
    assert invoked and result["review_complete"]
    assert result["files"][0]["render"]["status"] == "rendered"
    assert len(result["files"][0]["pages"]) == 2
    assert all(page["preview_url"] for page in result["files"][0]["pages"])


def test_active_pptx_content_is_not_sent_to_libreoffice(managed, monkeypatch):
    path = make_pptx(managed / "active.pptx")
    with zipfile.ZipFile(path, "a") as package:
        package.writestr("ppt/vbaProject.bin", b"not-executable-test-data")
    monkeypatch.setattr(review, "_find_libreoffice", lambda: pytest.fail("Renderer must not be invoked"))
    result = run(files=[str(path)])
    assert result["ok"] is False
    assert "活动内容" in result["error"]
