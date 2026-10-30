"""Long material content can be saved in batches and exported without resending it."""

from __future__ import annotations

import asyncio
import json
import zipfile
from pathlib import Path
from types import SimpleNamespace

import fitz
import pytest
from pptx import Presentation

from app.agent.tools import builtins, material_assets, material_workflow as workflow
from app.agent.tools.registry import execute_tool


@pytest.fixture
def authored(tmp_path, monkeypatch):
    exports = tmp_path / "exports"
    uploads = tmp_path / "uploads"
    workspace = tmp_path / "workspace"
    for folder in (exports, uploads, workspace):
        folder.mkdir()
    config = SimpleNamespace(
        resolve_path=lambda raw: exports if str(raw).endswith("/exports") else uploads,
        api_prefix="/api/v1",
    )
    monkeypatch.setattr(workflow, "exports_dir", lambda: exports)
    monkeypatch.setattr(builtins, "settings", config)
    monkeypatch.setattr(material_assets, "settings", config)
    monkeypatch.setattr("app.agent.tools.shell_config.resolve_shell", lambda: SimpleNamespace(
        enabled=False, workspace_path=str(workspace),
    ))
    return exports


def run(coro):
    return asyncio.run(coro)


def make_plan(*, pptx=0, pdf=False):
    deliverables = []
    outline = []
    if pptx:
        deliverables.append({"kind": "pptx", "purpose": "Classroom presentation", "total_slides": pptx})
        outline.extend({"kind": "pptx", "heading": f"Lesson {number}", "purpose": "Explain a concrete example"}
                       for number in range(1, pptx + 1))
    if pdf:
        deliverables.append({"kind": "pdf", "purpose": "Independent study"})
        outline.append({"kind": "pdf", "heading": "Handbook", "purpose": "Explain and practice"})
    return run(workflow.plan_materials(
        title="Training materials", audience="Java beginners", objectives=["Understand shared state"],
        deliverables=deliverables, outline=outline,
    ))


def write(saved, kind, items, draft_id=None):
    return run(workflow.write_material_draft(
        plan_id=saved["plan_id"], kind=kind, items=items, draft_id=draft_id,
    ))


def test_pptx_batches_merge_in_order_and_export_with_only_a_draft_reference(authored):
    saved = make_plan(pptx=8)
    slides = [{"layout": "title" if number == 1 else "bullets", "heading": f"Lesson {number}",
               "bullets": [f"Unique explanation for lesson {number}"]} for number in range(1, 9)]
    slides[3] = {"layout": "table", "heading": "Lesson 4", "columns": ["Thread", "Action"],
                 "rows": [["A", "Read zero"], ["B", "Read zero"]]}
    slides[4] = {"layout": "code", "heading": "Lesson 5", "language": "java",
                 "code": "int count = 0;\ncount++;", "explain": ["A compound operation"]}
    first = write(saved, "pptx", slides[:3])
    second = write(saved, "pptx", slides[3:6], first["draft_id"])
    final = write(saved, "pptx", slides[6:], second["draft_id"])
    assert final["item_count"] == final["expected_slide_count"] == 8
    assert "items" not in final  # Replies do not ask the model to resend old content.
    result = run(execute_tool("create_presentation", {
        "title": "Shared counter training", "draft_id": final["draft_id"],
    }))
    assert result["ok"], result
    assert result["plan_id"] == saved["plan_id"]
    presentation = Presentation(result["path"])
    assert len(presentation.slides) == 8
    for number, slide in enumerate(presentation.slides, 1):
        text = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
        assert f"Lesson {number}" in text
    assert any(shape.has_table for shape in presentation.slides[3].shapes)
    assert "count++;" in "\n".join(shape.text for shape in presentation.slides[4].shapes if shape.has_text_frame)


def test_long_pdf_chapters_export_without_retransmitting_the_content(authored):
    saved = make_plan(pdf=True)
    sections = [{
        "heading": f"Chapter {number}",
        "body": ("A shared counter needs a complete update protocol. "
                 "Readers should predict the outcome before looking at the answer. " * 20)
                + f" END_MARKER_{number}",
    } for number in range(1, 10)]
    first = write(saved, "pdf", sections[:3])
    middle = write(saved, "pdf", sections[3:6], first["draft_id"])
    last = write(saved, "pdf", sections[6:], middle["draft_id"])
    assert last["item_count"] == 9
    assert len(json.dumps(sections)) > 8192
    request = {"title": "Independent study handbook", "draft_id": last["draft_id"]}
    assert len(json.dumps(request)) < 200
    result = run(execute_tool("create_pdf", request))
    assert result["ok"], result
    assert result["plan_id"] == saved["plan_id"]
    with fitz.open(result["path"]) as document:
        assert len(document) > 3
        text = "\n".join(page.get_text() for page in document)
    for number in range(1, 10):
        assert f"END_MARKER_{number}" in text


def test_source_batches_keep_original_compilable_filenames_and_content(authored):
    saved = make_plan(pdf=True)
    code = 'public class CounterLab {\n    public static void main(String[] args) {\n        System.out.println("done");\n    }\n}\n'
    readme = "Compile: javac CounterLab.java\nRun: java CounterLab\n"
    first = write(saved, "sources", [{"filename": "CounterLab.java", "content": code}])
    last = write(saved, "sources", [{"filename": "README.md", "content": readme}], first["draft_id"])
    result = run(execute_tool("create_code_files", {
        "title": "Counter laboratory", "draft_id": last["draft_id"],
    }))
    assert result["ok"], result
    assert result["verification_status"] == "not_run"
    with zipfile.ZipFile(result["path"]) as archive:
        assert archive.namelist() == ["CounterLab.java", "README.md"]
        assert archive.read("CounterLab.java").decode("utf-8") == code
        assert archive.read("README.md").decode("utf-8") == readme


def test_draft_cannot_cross_plans_or_formats(authored):
    saved = make_plan(pptx=2, pdf=True)
    other = make_plan(pptx=2, pdf=True)
    first = write(saved, "pptx", [{"layout": "title", "heading": "Lesson 1"}])
    original = (authored / f"material_draft_{first['draft_id']}.json").read_bytes()
    with pytest.raises(ValueError, match="format/plan"):
        workflow.load_material_draft(first["draft_id"], "pdf", saved["plan_id"])
    with pytest.raises(ValueError, match="format/plan"):
        workflow.load_material_draft(first["draft_id"], "pptx", other["plan_id"])
    with pytest.raises(ValueError, match="format/plan"):
        write(other, "pptx", [{"layout": "bullets", "heading": "Wrong plan"}], first["draft_id"])
    mismatch = run(execute_tool("create_pdf", {"title": "Wrong format", "draft_id": first["draft_id"]}))
    assert mismatch["ok"] is False
    assert (authored / f"material_draft_{first['draft_id']}.json").read_bytes() == original


def test_format_not_in_plan_is_rejected_before_creating_a_draft(authored):
    saved = make_plan(pdf=True)
    with pytest.raises(ValueError, match="format"):
        write(saved, "pptx", [{"heading": "Unexpected slides"}])
    assert not list(authored.glob("material_draft_*"))


def test_exceeding_planned_slide_total_does_not_discard_existing_batches(authored):
    saved = make_plan(pptx=3)
    slides = [{"layout": "title" if number == 1 else "bullets", "heading": f"Lesson {number}"}
              for number in range(1, 4)]
    current = write(saved, "pptx", slides)
    path = authored / f"material_draft_{current['draft_id']}.json"
    original = path.read_bytes()
    with pytest.raises(ValueError, match="exceeding"):
        write(saved, "pptx", [{"heading": "Unrequested fourth slide"}], current["draft_id"])
    assert path.read_bytes() == original
    assert workflow.load_material_draft(current["draft_id"], "pptx")["items"] == slides
    assert not list(authored.glob("*.tmp"))


def test_incomplete_pptx_draft_can_be_completed_after_export_rejects_it(authored):
    saved = make_plan(pptx=2)
    first = write(saved, "pptx", [{"layout": "title", "heading": "Lesson 1"}])
    result = run(execute_tool("create_presentation", {"title": "Not finished", "draft_id": first["draft_id"]}))
    assert result["ok"] is False
    assert len(workflow.load_material_draft(first["draft_id"], "pptx")["items"]) == 1
    write(saved, "pptx", [{"layout": "bullets", "heading": "Lesson 2"}], first["draft_id"])
    result = run(execute_tool("create_presentation", {"title": "Now complete", "draft_id": first["draft_id"]}))
    assert result["ok"] and result["slide_count"] == 2


def test_draft_and_inline_content_cannot_be_supplied_together(authored):
    saved = make_plan(pptx=1, pdf=True)
    deck = write(saved, "pptx", [{"layout": "title", "heading": "Lesson 1"}])
    handbook = write(saved, "pdf", [{"heading": "Overview", "body": "An explanation."}])
    for name, payload in (
        ("create_presentation", {"draft_id": deck["draft_id"], "slides": [{"heading": "Conflicting inline slide"}]}),
        ("create_pdf", {"draft_id": handbook["draft_id"], "sections": [{"heading": "Conflicting inline chapter"}]}),
    ):
        result = run(execute_tool(name, {"title": "Conflicting request", **payload}))
        assert result["ok"] is False
        assert "not both" in result["error"]


def test_material_loop_budget_does_not_change_regular_chat_limits(monkeypatch):
    from app.agent.nodes.llm_call import should_call_tools
    from app.agent.state import TOOL_LOOP_GRACE, tool_loop_limit
    from app.core.config import settings

    monkeypatch.setattr(settings, "agent_max_tool_loops", 12)
    monkeypatch.setattr(settings, "materials_max_tool_loops", 24)
    tool = [{"function": {"name": "web_search"}}]
    ordinary = {"required_tools": ["web_search"], "tool_calls": tool}
    material = {"required_tools": ["plan_materials", "create_pdf"], "tool_calls": tool}
    assert tool_loop_limit(ordinary) == 12
    assert tool_loop_limit(material) == 24
    assert should_call_tools({**ordinary, "tool_loop_count": 12 + TOOL_LOOP_GRACE}) == "verify"
    assert should_call_tools({**material, "tool_loop_count": 12 + TOOL_LOOP_GRACE}) == "tools"
    assert should_call_tools({**material, "tool_loop_count": 24 + TOOL_LOOP_GRACE}) == "verify"


def test_longer_model_timeout_applies_only_to_material_workflow(monkeypatch):
    from app.agent.nodes import llm_call
    from app.core.config import settings
    from app.llm.client import CompletionResult
    from app.llm.providers import ResolvedLLM

    monkeypatch.setattr(settings, "materials_llm_timeout_seconds", 300.0)
    configured = ResolvedLLM(
        source="runtime", kind="openai_compatible", model="test-model", label="Test model",
        temperature=0.0, max_tokens=8192, timeout_seconds=120.0, max_retries=0,
    )
    calls = []

    async def schemas(**kwargs):
        return []

    async def model(messages, **kwargs):
        calls.append(kwargs)
        return CompletionResult(content="Ready")

    monkeypatch.setattr(llm_call, "openai_tool_schemas", schemas)
    monkeypatch.setattr(llm_call, "chat_complete", model)
    monkeypatch.setattr(llm_call, "has_sink", lambda: False)
    monkeypatch.setattr("app.llm.client._config_for_call", lambda override: configured)
    run(llm_call.llm_call_node({"messages": [], "required_tools": []}, session=object()))
    run(llm_call.llm_call_node({"messages": [], "required_tools": ["plan_materials"]}, session=object()))
    assert "config" not in calls[0]
    assert calls[1]["config"].timeout_seconds == 300.0
    assert configured.timeout_seconds == 120.0


def test_bounded_draft_read_keeps_full_chapter_content_after_tool_execution(authored, monkeypatch):
    from app.agent.nodes.tool_executor import tool_executor_node
    from app.core.config import settings

    saved = make_plan(pdf=True)
    body = "并发知识" * 5000 + " UNIQUE_MIDDLE_MARKER " + "并发知识" * 5000
    current = write(saved, "pdf", [{"heading": "Complete chapter", "body": body}])
    monkeypatch.setattr(settings, "spill_max_inline_bytes", 500)
    monkeypatch.setattr(settings, "spill_dir", str(authored.parent / "spills"))
    state = run(tool_executor_node({
        "user_id": "test-user", "conversation_id": "recover-chapter", "messages": [],
        "required_tools": ["create_pdf", "review_materials"],
        "tool_calls": [{"id": "read-chapter", "function": {"name": "read_material_draft", "arguments": json.dumps({
            "draft_id": current["draft_id"], "kind": "pdf", "start": 0, "count": 1,
        })}}],
    }, session=object()))
    response = json.loads(state["messages"][-1]["content"])
    assert response["ok"]
    assert not response.get("spilled")
    assert response["items"][0]["body"] == body
    assert "UNIQUE_MIDDLE_MARKER" in response["items"][0]["body"]


def test_draft_read_byte_limit_can_be_recovered_with_smaller_ranges(authored):
    saved = make_plan(pdf=True)
    chapter = {"heading": "A detailed chapter", "body": "知识" * 11000}
    current = write(saved, "pdf", [chapter, chapter])
    too_large = run(execute_tool("read_material_draft", {
        "draft_id": current["draft_id"], "kind": "pdf", "start": 0, "count": 2,
    }))
    assert too_large["ok"] is False
    assert "128000" in too_large["error"]
    smaller = run(execute_tool("read_material_draft", {
        "draft_id": current["draft_id"], "kind": "pdf", "start": 1, "count": 1,
    }))
    assert smaller["ok"] and smaller["item_count"] == 2
    assert smaller["items"] == [chapter]
    assert len(workflow.load_material_draft(current["draft_id"], "pdf")["items"]) == 2


def test_saved_draft_metadata_survives_trace_and_pure_continue_restores_it(authored):
    from app.agent.execution_trace import tool_step
    from app.agent.nodes.compose_prompt import compose_prompt_node
    from app.agent.nodes.load_memory import material_context
    from app.agent.nodes.plan import plan_node
    from app.agent.nodes.tool_executor import tool_executor_node
    from app.agent.state import todo_only_loop_limit, tool_loop_limit

    saved = make_plan(pptx=8, pdf=True)
    saved_trace = tool_step(name="plan_materials", call_id="plan", raw_args={}, result=saved)
    written = run(tool_executor_node({
        "user_id": "test-user", "conversation_id": "save-draft", "messages": [],
        "material_plan_id": saved["plan_id"],
        "tool_calls": [{"id": "write-chapter", "function": {"name": "write_material_draft", "arguments": json.dumps({
            "kind": "pdf", "items": [{"heading": "Saved chapter", "body": "Explain shared mutable state."}],
        })}}],
    }, session=object()))
    draft_step = written["execution_trace"][-1]
    metadata = draft_step["material_draft"]
    assert metadata["kind"] == "pdf" and metadata["item_count"] == 1
    restored = material_context([SimpleNamespace(execution_trace=[saved_trace, draft_step])])
    assert restored["material_drafts"]["pdf"] == metadata
    state = run(plan_node({
        "query": "继续", "conversation_id": "resume-draft", **restored,
        "history": [{"role": "user", "content": "制作8页PPT与PDF讲义"}],
    }))
    assert state["material_resume"] is True
    assert "plan_materials" not in state["required_tools"]
    assert {"create_presentation", "create_pdf", "review_materials"} <= set(state["required_tools"])
    assert todo_only_loop_limit(state) == 12
    assert tool_loop_limit(state) >= 24
    prompted = run(compose_prompt_node(state))
    system = prompted["messages"][0]["content"]
    assert metadata["draft_id"] in system
    assert "read_material_draft" in system
    assert "无需重新调用plan_materials" in system
    changed = run(plan_node({**restored, "query": "继续，把PPT改成6页", "history": []}))
    assert changed["material_resume"] is False
    assert "plan_materials" in changed["required_tools"]


def test_restoring_new_plan_drops_drafts_from_old_plan(authored):
    from app.agent.execution_trace import tool_step
    from app.agent.nodes.load_memory import material_context

    first = make_plan(pdf=True)
    second = make_plan(pdf=True)
    old_draft = write(first, "pdf", [{"heading": "Old lesson", "body": "Previous content"}])
    new_draft = write(second, "pdf", [{"heading": "New lesson", "body": "Current content"}])
    steps = [
        tool_step(name="plan_materials", call_id="old-plan", raw_args={}, result=first),
        tool_step(name="write_material_draft", call_id="old-draft", raw_args={}, result=old_draft),
        tool_step(name="plan_materials", call_id="new-plan", raw_args={}, result=second),
        tool_step(name="write_material_draft", call_id="new-draft", raw_args={}, result=new_draft),
        tool_step(name="write_material_draft", call_id="stale-draft", raw_args={}, result=old_draft),
    ]
    restored = material_context([SimpleNamespace(execution_trace=steps)])
    assert restored["material_plan_id"] == second["plan_id"]
    assert restored["material_drafts"]["pdf"]["draft_id"] == new_draft["draft_id"]


def test_fourth_checklist_update_does_not_stop_a_multistep_material_turn(authored):
    from app.agent.nodes.llm_call import should_call_tools
    from app.agent.nodes.tool_executor import tool_executor_node

    saved = make_plan(pdf=True)
    state = {
        "user_id": "test-user", "conversation_id": "material-checkpoints", "messages": [],
        "material_plan_id": saved["plan_id"],
        "required_tools": ["plan_materials", "create_pdf", "review_materials"],
    }
    for checkpoint in range(4):
        state["tool_calls"] = [{"id": f"check-{checkpoint}", "function": {
            "name": "todo_write", "arguments": json.dumps({"todos": [
                {"content": f"Complete material stage {checkpoint + 1}", "status": "in_progress"}
            ]}),
        }}]
        assert should_call_tools(state) == "tools"
        state = run(tool_executor_node(state, session=object()))
    assert state["todo_loop_count"] == 4
    assert state["tool_loop_count"] == 0
    state["tool_calls"] = [{"id": "real-work", "function": {"name": "write_material_draft", "arguments": json.dumps({
        "kind": "pdf", "items": [{"heading": "A real chapter", "body": "Saved after four checklist updates."}],
    })}}]
    assert should_call_tools(state) == "tools"
    completed = run(tool_executor_node(state, session=object()))
    assert completed["tool_loop_count"] == 1
    assert completed["todo_loop_count"] == 0
    assert completed["execution_trace"][-1]["material_draft"]["item_count"] == 1


def test_regular_chat_checklist_budget_remains_three():
    from app.agent.nodes.llm_call import should_call_tools
    from app.agent.state import todo_only_loop_limit

    call = [{"function": {"name": "todo_write"}}]
    ordinary = {"required_tools": ["web_search"], "tool_calls": call}
    assert todo_only_loop_limit(ordinary) == 3
    assert should_call_tools({**ordinary, "todo_loop_count": 2}) == "tools"
    assert should_call_tools({**ordinary, "todo_loop_count": 3}) == "verify"
    material = {"required_tools": ["review_materials"], "tool_calls": call}
    assert should_call_tools({**material, "todo_loop_count": 11}) == "tools"
    assert should_call_tools({**material, "todo_loop_count": 12}) == "verify"


def test_exhausted_retries_do_not_mean_the_delivery_was_verified(monkeypatch):
    from app.agent.nodes.verify import verify_node
    from app.core.config import settings

    monkeypatch.setattr(settings, "agent_max_tool_loops", 12)
    monkeypatch.setattr(settings, "materials_max_tool_loops", 24)
    failed = run(verify_node({
        "required_tools": ["create_pdf", "review_materials"], "tool_loop_count": 26,
        "response": "Still preparing the requested handout.", "messages": [],
    }))
    assert failed["needs_retry"] is False
    assert failed["verification_complete"] is False
    passed = run(verify_node({"required_tools": [], "response": "Hello.", "messages": []}))
    assert passed["needs_retry"] is False
    assert passed["verification_complete"] is True
