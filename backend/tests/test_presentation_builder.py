"""Material authoring contracts: editable evidence, exact counts and safe assets."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.agent.tools import material_assets
from app.agent.tools.pptx_builder import build_presentation


@pytest.fixture
def material_dirs(tmp_path: Path, monkeypatch):
    for name in ("exports", "uploads"):
        (tmp_path / "data" / name).mkdir(parents=True)
    monkeypatch.setattr(
        material_assets, "settings",
        SimpleNamespace(api_prefix="/api/v1", resolve_path=lambda raw: tmp_path / raw.removeprefix("./")),
    )
    return tmp_path


def _roundtrip(presentation):
    stream = BytesIO()
    presentation.save(stream)
    stream.seek(0)
    return Presentation(stream)


def _texts(slide):
    return [shape.text for shape in slide.shapes if shape.has_text_frame]


def test_excess_points_and_cards_are_rejected_instead_of_discarded():
    with pytest.raises(ValueError, match="10 bullet"):
        build_presentation(title="Long list", expected_slide_count=1,
                           slides=[{"heading": "All required points", "bullets": [str(i) for i in range(11)]}])
    with pytest.raises(ValueError, match="4 items"):
        build_presentation(title="Many cards", expected_slide_count=1,
                           slides=[{"layout": "cards", "heading": "All required cards", "cards": [{"label": str(i), "value": str(i)} for i in range(5)]}])


def test_exact_count_uses_only_supplied_slides_and_real_page_numbers():
    deck = _roundtrip(build_presentation(
        title="培训", expected_slide_count=2,
        slides=[{"layout": "title", "heading": "培训"},
                {"heading": "一个明确结论", "bullets": ["解释这个结论。"]}],
    ))
    assert len(deck.slides) == 2
    assert "2 / 2" in _texts(deck.slides[1])
    assert not any("谢谢" in text for slide in deck.slides for text in _texts(slide))


def test_exact_count_without_cover_does_not_insert_one():
    deck = build_presentation(title="单页说明", expected_slide_count=1,
                              slides=[{"heading": "直接讲内容", "bullets": ["第一点"]}])
    assert len(deck.slides) == 1


@pytest.mark.parametrize("expected", [0, 61, True, 2.5])
def test_invalid_exact_count_is_rejected(expected):
    with pytest.raises(ValueError, match="expected_slide_count"):
        build_presentation(title="示例", slides=[{"heading": "正文"}], expected_slide_count=expected)


def test_exact_count_mismatch_is_actionable():
    with pytest.raises(ValueError, match="Expected exactly 8 slides.*received 1"):
        build_presentation(title="示例", slides=[{"heading": "正文"}], expected_slide_count=8)


def test_legacy_calls_still_get_automatic_cover():
    deck = build_presentation(title="旧调用", slides=[{"heading": "正文", "bullets": ["内容"]}])
    assert len(deck.slides) == 2
    assert "旧调用" in _texts(deck.slides[0])


def test_table_is_native_editable_and_retains_values_and_sources():
    deck = _roundtrip(build_presentation(
        title="执行过程", expected_slide_count=1,
        slides=[{
            "layout": "table", "heading": "交错推演",
            "columns": ["步骤", "线程 A", "线程 B", "共享值"],
            "rows": [[1, "读取 0", "等待", 0], [2, "等待", "读取 0", 0],
                     [3, "写回 1", "等待", 1], [4, "结束", "写回 1", 1]],
            "caption": "教学调度示意",
            "speaker_notes": ["先让学员预测。", "按行推演每个线程的局部值。"],
            "references": [{"title": "规范", "url": "https://example.org/spec"}],
        }],
    ))
    tables = [shape.table for shape in deck.slides[0].shapes if shape.has_table]
    assert len(tables) == 1
    assert tables[0].cell(4, 2).text == "写回 1"
    assert tables[0].cell(4, 3).text == "1"
    tables[0].cell(4, 3).text = "可编辑"
    notes = deck.slides[0].notes_slide.notes_text_frame.text
    assert "先让学员预测" in notes
    assert "规范: https://example.org/spec" in notes


def test_table_rejects_lost_rows_ragged_rows_and_dense_cells():
    common = {"layout": "table", "heading": "数据", "columns": ["名称", "值"]}
    with pytest.raises(ValueError, match="9 data rows"):
        build_presentation(title="表格", expected_slide_count=1,
                           slides=[{**common, "rows": [[i, i] for i in range(9)]}])
    with pytest.raises(ValueError, match="2 cells"):
        build_presentation(title="表格", expected_slide_count=1,
                           slides=[{**common, "rows": [[1]]}])
    with pytest.raises(ValueError, match="too dense"):
        build_presentation(title="表格", expected_slide_count=1,
                           slides=[{**common, "rows": [["长" * 1000, "值"]]}])


def test_code_preserves_spaces_and_last_line_without_line_numbers():
    code = "public class Demo {\n    public static void main(String[] args) {\n        System.out.println(1);\n    }\n}"
    deck = _roundtrip(build_presentation(
        title="代码", expected_slide_count=1,
        slides=[{"layout": "code", "heading": "完整示例", "code": code, "language": "java"}],
    ))
    sources = [text for text in _texts(deck.slides[0]) if "public class Demo" in text]
    assert sources == [code]
    assert "\u00a0" not in sources[0]


def test_code_has_no_silent_line_or_width_truncation():
    code = "\n".join(f"line{i}();" for i in range(24))
    deck = build_presentation(title="代码", expected_slide_count=1,
                              slides=[{"layout": "code", "heading": "边界", "code": code}])
    assert any(text.endswith("line23();") for text in _texts(deck.slides[0]))
    with pytest.raises(ValueError, match="25 lines.*limit is 24"):
        build_presentation(title="代码", expected_slide_count=1,
                           slides=[{"layout": "code", "heading": "太长", "code": code + "\nlast();"}])
    with pytest.raises(ValueError, match="too wide"):
        build_presentation(title="代码", expected_slide_count=1,
                           slides=[{"layout": "code", "heading": "太宽", "code": "x" * 220}])
    with pytest.raises(ValueError, match="requires real code"):
        build_presentation(title="代码", expected_slide_count=1,
                           slides=[{"layout": "code", "heading": "缺少源码"}])


def test_cover_and_figure_embed_existing_images_and_keep_aspect(material_dirs):
    image = material_dirs / "data" / "uploads" / "source photo.png"
    Image.new("RGB", (200, 100), "teal").save(image)
    source = "/api/v1/files/uploads/source%20photo.png"
    deck = _roundtrip(build_presentation(
        title="带图课程", expected_slide_count=2, cover_image=source,
        slides=[{"layout": "title", "heading": "课程"},
                {"layout": "figure", "heading": "原始图", "image": source, "caption": "原图说明"}],
    ))
    pictures = [[shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
                for slide in deck.slides]
    assert [len(items) for items in pictures] == [1, 1]
    assert abs(pictures[1][0].width / pictures[1][0].height - 2) < 0.01
    assert "原图说明" in _texts(deck.slides[1])


def test_figure_missing_image_is_an_error(material_dirs):
    with pytest.raises(ValueError, match="requires an image"):
        build_presentation(title="图", expected_slide_count=1,
                           slides=[{"layout": "figure", "heading": "必须有图"}])
    with pytest.raises(ValueError, match="does not exist"):
        build_presentation(title="图", expected_slide_count=1,
                           slides=[{"layout": "figure", "heading": "缺图", "image": "uploads/missing.png"}])


def test_asset_resolver_accepts_local_files_and_custom_api_prefix(material_dirs, monkeypatch):
    document = material_dirs / "data" / "exports" / "lesson.pdf"
    document.write_bytes(b"not an image; the resolver is format-agnostic")
    assert material_assets.resolve_material_asset(str(document)) == document.resolve()
    assert material_assets.resolve_material_asset("exports/lesson.pdf") == document.resolve()
    assert material_assets.resolve_material_asset("/api/v1/files/exports/lesson.pdf") == document.resolve()
    monkeypatch.setattr(material_assets.settings, "api_prefix", "/custom")
    assert material_assets.resolve_material_asset("/custom/files/exports/lesson.pdf") == document.resolve()


@pytest.mark.parametrize("source", [
    "https://example.org/photo.png", "http://127.0.0.1/private.png",
    "file:///etc/passwd", "//example.org/image.png", "exports/../secret.txt",
    "/api/v1/files/exports/%2e%2e/secret.txt", "uploads/%2e%2e%5csecret.txt",
    "uploads/image.png?token=secret", "uploads/image.png#fragment", "uploads/a\x00.png",
])
def test_asset_resolver_rejects_remote_paths_and_traversal(material_dirs, source):
    with pytest.raises(ValueError):
        material_assets.resolve_material_asset(source)


def test_asset_resolver_rejects_absolute_server_files_and_symlink_escape(material_dirs):
    secret = material_dirs / "secret.txt"
    secret.write_text("private server file", encoding="utf-8")
    with pytest.raises(ValueError, match="inside data/exports or data/uploads"):
        material_assets.resolve_material_asset(str(secret))
    link = material_dirs / "data" / "exports" / "shortcut.txt"
    try:
        link.symlink_to(secret)
    except OSError:
        pytest.skip("Creating symlinks requires additional privileges on this host.")
    with pytest.raises(ValueError, match="outside its managed directory"):
        material_assets.resolve_material_asset("exports/shortcut.txt")
