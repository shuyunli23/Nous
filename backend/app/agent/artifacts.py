"""Read back files the agent just produced — Codex-style observe, not vibes.

The model is not allowed to claim a webpage/PPT/PDF contains something
unless this inspector saw it in the bytes on disk.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any
from urllib.parse import unquote

from app.agent.tools.export_urls import EXPORT_URL_RE

_TAG = re.compile(r"<[^>]+>")
_IMG = re.compile(r"<img\b[^>]*>", re.I)
_H = re.compile(r"<h[123][^>]*>(.*?)</h[123]>", re.I | re.S)
_SECTION = re.compile(r"<section\b[^>]*>(.*?)</section>", re.I | re.S)
_PRE_BLOCK = re.compile(r"<pre\b[^>]*>(.*?)</pre>", re.I | re.S)
_DONE = re.compile(
    r"(已经|已)(成功)?(更新|添加|加上|完成|制作|放入|嵌|加好|改好)"
)
_VISUAL = re.compile(
    r"(照片|图像|头像|一寸|证件照|图片|配图|插图|figure|portrait|headshot)"
)
_CODE_CLAIM = re.compile(r"伪代码|伪码|pseudocode|算法代码")
_CODE_HEADING = re.compile(r"伪代码|伪码|pseudocode|algorithm", re.I)
_ADD = ("加", "加上", "放", "截", "裁", "贴", "换", "补", "嵌")
_SOURCE_VIEW = ("源码", "源代码", "source code", "查看源码", "给我源码")


def inspect_path(path: Path | str | None) -> dict[str, Any] | None:
    if not path:
        return None
    file = Path(path)
    if not file.is_file():
        return None
    suffix = file.suffix.lower()
    if suffix in {".html", ".htm"}:
        return _inspect_html(file)
    if suffix == ".pptx":
        return _inspect_pptx(file)
    if suffix == ".pdf":
        return _inspect_pdf(file)
    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".gif"}:
        return _inspect_image(file)
    return {
        "kind": "file",
        "file": file.name,
        "bytes": file.stat().st_size,
    }


def attach_inspect(result: dict[str, Any]) -> dict[str, Any]:
    """Mutate a tool result so the LLM sees what is actually on disk."""
    if not result.get("ok"):
        return result
    raw = result.get("path")
    inspect = inspect_path(raw) if raw else None
    if inspect is None:
        url = str(result.get("download_url") or "")
        name = unquote(url.rsplit("/", 1)[-1].split("?", 1)[0]) if url else ""
        if name:
            from app.core.config import settings

            inspect = inspect_path(settings.resolve_path("./data/exports") / name)
    if inspect:
        result["inspect"] = inspect
    return result


def inspects_from_trace(trace: list[dict[str, Any]] | None) -> list[dict[str, Any]]:
    found: list[dict[str, Any]] = []
    for step in trace or []:
        if step.get("kind") != "tool" or step.get("status") == "error":
            continue
        inspect = step.get("inspect")
        if isinstance(inspect, dict):
            found.append(inspect)
            continue
        detail = str(step.get("detail") or "")
        for url in EXPORT_URL_RE.findall(detail):
            name = unquote(url.rsplit("/", 1)[-1].split("?", 1)[0])
            if not name:
                continue
            from app.core.config import settings

            inspect = inspect_path(settings.resolve_path("./data/exports") / name)
            if inspect:
                found.append(inspect)
    return found


def latest_inspect(
    inspects: list[dict[str, Any]], kind: str
) -> dict[str, Any] | None:
    for item in reversed(inspects):
        if item.get("kind") == kind:
            return item
    return None


def visual_requested(goal: str) -> bool:
    text = (goal or "").strip()
    if not text:
        return False
    if any(key in text for key in ("一寸", "证件照", "头像")):
        return True
    return bool(_VISUAL.search(text)) and any(key in text for key in _ADD)


def code_requested(goal: str) -> bool:
    """User asked to put pseudocode / algorithm code into the demo page."""
    text = (goal or "").strip()
    if not text:
        return False
    if _CODE_CLAIM.search(text):
        return True
    if any(key in text.lower() for key in _SOURCE_VIEW) and "伪代码" not in text:
        return False
    if "算法" in text and any(key in text for key in _ADD + ("写", "展示", "加入")):
        return True
    return False


def delivery_gaps(
    *,
    goal: str,
    response: str,
    inspects: list[dict[str, Any]],
) -> list[str]:
    """Ground-truth mismatches between what was asked/claimed and the files."""
    gaps: list[str] = []
    web = latest_inspect(inspects, "webpage")
    ppt = latest_inspect(inspects, "pptx")
    claimed_visual = bool(
        _DONE.search(response or "") and _VISUAL.search(response or "")
    )
    wants_visual = visual_requested(goal) or claimed_visual
    image_count = 0
    hero_photo = False
    if web:
        image_count += int(web.get("image_count") or 0)
        hero_photo = bool(web.get("hero_has_photo"))
    if ppt:
        image_count += int(ppt.get("image_count") or 0)
    if wants_visual and image_count <= 0:
        gaps.append(
            "inspect.image_count=0 — the file has no <img>/picture. "
            "Do not say the photo was added. Put a real upload/crop URL on "
            "hero.image (or a PPT picture) and rebuild."
        )
    elif visual_requested(goal) and web and not hero_photo and image_count > 0:
        gaps.append(
            "inspect.hero_has_photo=false — images exist but the cover has no "
            "portrait. Pass the photo as hero.image and call create_webpage again."
        )

    claimed_code = bool(
        _DONE.search(response or "") and _CODE_CLAIM.search(response or "")
    )
    wants_code = code_requested(goal) or claimed_code
    empty_code_headings = [
        str(h)
        for h in (web.get("empty_headings") if web else []) or []
        if _CODE_HEADING.search(str(h) or "")
    ]
    has_code = bool(web and web.get("has_code"))
    if web and (wants_code or empty_code_headings) and not has_code:
        gaps.append(
            "inspect.has_code=false — the page has no real <pre> code block. "
            "A heading like 伪代码 with an empty body does not count. "
            "Call create_webpage again with layout=code and a multiline `code` field."
        )
    elif web and empty_code_headings:
        gaps.append(
            "inspect.empty_headings="
            + " | ".join(empty_code_headings)
            + " — that section is title-only. Fill `code` and rebuild."
        )
    return gaps


def format_inspect_for_prompt(inspects: list[dict[str, Any]]) -> str:
    if not inspects:
        return "inspect: (no files produced this turn)"
    chunks = []
    for item in inspects[-3:]:
        kind = item.get("kind") or "file"
        name = item.get("file") or ""
        bits = [f"{kind}:{name}"]
        if item.get("section_count") is not None:
            bits.append(f"sections={item.get('section_count')}")
        if item.get("slide_count") is not None:
            bits.append(f"slides={item.get('slide_count')}")
        if item.get("page_count") is not None:
            bits.append(f"pages={item.get('page_count')}")
        if item.get("image_count") is not None:
            bits.append(f"images={item.get('image_count')}")
        if "hero_has_photo" in item:
            bits.append(f"hero_photo={item.get('hero_has_photo')}")
        if item.get("has_code") is not None:
            bits.append(f"has_code={item.get('has_code')}")
        if item.get("code_chars"):
            bits.append(f"code_chars={item.get('code_chars')}")
        empty = item.get("empty_headings") or []
        if empty:
            bits.append("empty=" + " | ".join(str(h) for h in empty[:6]))
        headings = item.get("headings") or item.get("titles") or []
        if headings:
            bits.append("headings=" + " | ".join(str(h) for h in headings[:8]))
        chunks.append("; ".join(bits))
    return "inspect: " + " || ".join(chunks)


def _plain(html: str) -> str:
    return re.sub(r"\s+", " ", _TAG.sub("", html)).strip()


def _inspect_html(path: Path) -> dict[str, Any]:
    html = path.read_text(encoding="utf-8", errors="ignore")
    headings = [_plain(h) for h in _H.findall(html)]
    headings = [h for h in headings if h][:12]
    empty_headings: list[str] = []
    for raw in _SECTION.findall(html):
        heads = [_plain(h) for h in _H.findall(raw)]
        heading = next((h for h in heads if h), "")
        leftover = _plain(raw)
        if heading:
            leftover = leftover.replace(heading, "", 1).strip()
        if heading and len(leftover) < 12:
            empty_headings.append(heading)
    code_chars = sum(len(_plain(block)) for block in _PRE_BLOCK.findall(html))
    return {
        "kind": "webpage",
        "file": path.name,
        "bytes": path.stat().st_size,
        "section_count": html.lower().count("<section"),
        "image_count": len(_IMG.findall(html)),
        "hero_has_photo": "hero__photo" in html,
        "has_code": code_chars >= 40,
        "code_chars": code_chars,
        "empty_headings": empty_headings[:8],
        "headings": headings,
    }


def _inspect_pptx(path: Path) -> dict[str, Any]:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:
        return {"kind": "pptx", "file": path.name, "bytes": path.stat().st_size}
    prs = Presentation(str(path))
    titles: list[str] = []
    pictures = 0
    for slide in prs.slides:
        title = ""
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and not title:
                text = (shape.text_frame.text or "").strip()
                if text:
                    title = text.splitlines()[0][:80]
            try:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    pictures += 1
            except Exception:
                continue
        if title:
            titles.append(title)
    return {
        "kind": "pptx",
        "file": path.name,
        "bytes": path.stat().st_size,
        "slide_count": len(prs.slides),
        "image_count": pictures,
        "titles": titles[:12],
    }


def _inspect_pdf(path: Path) -> dict[str, Any]:
    page_count = 0
    preview = ""
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        page_count = len(reader.pages)
        bits = []
        for page in reader.pages[:2]:
            bits.append((page.extract_text() or "").strip())
        preview = " ".join(bits)[:240]
    except Exception:
        page_count = 0
    return {
        "kind": "pdf",
        "file": path.name,
        "bytes": path.stat().st_size,
        "page_count": page_count,
        "preview": preview,
    }


def _inspect_image(path: Path) -> dict[str, Any]:
    info: dict[str, Any] = {
        "kind": "image",
        "file": path.name,
        "bytes": path.stat().st_size,
        "image_count": 1,
    }
    try:
        from PIL import Image

        with Image.open(path) as image:
            info["width"] = image.size[0]
            info["height"] = image.size[1]
    except Exception:
        pass
    return info
