"""Save chat uploads, extract document text, and prepare vision parts."""

from __future__ import annotations

import base64
import csv
import json
import re
import uuid
from dataclasses import dataclass, field
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any, Literal

from app.core.config import settings
from app.core.exceptions import ValidationError
from app.core.logging import get_logger

logger = get_logger(__name__)

ATTACH_START = "<!--nous-attachments-->"
ATTACH_END = "<!--/nous-attachments-->"

MAX_FILES = 8
MAX_FILE_BYTES = 100 * 1024 * 1024
MAX_EXTRACT_PER_FILE = 14_000
MAX_EXTRACT_TOTAL = 26_000
MAX_QUERY_CHARS = 30_000
MAX_VISION_IMAGES = 4
MAX_VISION_BYTES = 4 * 1024 * 1024
MAX_PDF_FIGURES = 8

Kind = Literal["document", "image", "unsupported"]

_SAFE_STEM = re.compile(r"[^\w\u4e00-\u9fff.\-]+")

DOCUMENT_EXTS = {
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".log",
    ".xml",
    ".html",
    ".htm",
    ".pdf",
    ".docx",
    ".doc",
    ".pptx",
    ".ppt",
    ".xlsx",
    ".xlsm",
    ".xls",
}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
ALLOWED_EXTS = DOCUMENT_EXTS | IMAGE_EXTS

EXT_MIME = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".csv": "text/csv",
    ".json": "application/json",
    ".log": "text/plain",
    ".xml": "application/xml",
    ".html": "text/html",
    ".htm": "text/html",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".doc": "application/msword",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".ppt": "application/vnd.ms-powerpoint",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12",
    ".xls": "application/vnd.ms-excel",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".gif": "image/gif",
    ".bmp": "image/bmp",
}

_VISION_MODEL_HINTS = (
    "claude-3",
    "claude-4",
    "claude-sonnet",
    "claude-haiku",
    "claude-opus",
    "sonnet",
    "haiku",
    "opus",
    "gpt-4o",
    "gpt-4.1",
    "gpt-4-turbo",
    "gpt-5",
    "gemini",
    "nova",
    "qwen-vl",
    "qwen2.5-vl",
    "qwen2-vl",
    "llava",
    "pixtral",
    "grok-2-vision",
    "vision",
)


@dataclass
class ProcessedUpload:
    original_name: str
    stored_name: str
    download_url: str
    kind: Kind
    mime: str
    extract: str = ""
    vision: dict[str, Any] | None = None
    extra_visions: list[dict[str, Any]] = field(default_factory=list)
    figures: list[dict[str, Any]] = field(default_factory=list)
    error: str | None = None
    notes: list[str] = field(default_factory=list)


def process_chat_uploads(
    files: list[tuple[str, bytes, str]],
) -> list[ProcessedUpload]:
    """Validate, persist, and extract each uploaded file."""
    if len(files) > MAX_FILES:
        raise ValidationError(
            f"Too many files (max {MAX_FILES}).",
            details={"max_files": MAX_FILES},
        )

    upload_dir = settings.resolve_path("./data/uploads")
    upload_dir.mkdir(parents=True, exist_ok=True)

    processed: list[ProcessedUpload] = []
    vision_count = 0
    for original_name, data, content_type in files:
        name = (original_name or "unnamed").strip() or "unnamed"
        if len(data) > MAX_FILE_BYTES:
            raise ValidationError(
                f"File '{name}' is too large (max {MAX_FILE_BYTES // (1024 * 1024)} MB).",
                details={"filename": name, "max_bytes": MAX_FILE_BYTES},
            )
        if not data:
            raise ValidationError(
                f"File '{name}' is empty.",
                details={"filename": name},
            )

        ext = Path(name).suffix.lower()
        if ext not in ALLOWED_EXTS:
            raise ValidationError(
                f"Unsupported file type '{ext or name}'. "
                "Use txt, md, csv, json, pdf, Word, PPT, Excel, or a common image.",
                details={"filename": name, "extension": ext},
            )

        mime = (content_type or "").split(";")[0].strip().lower()
        if not mime or mime == "application/octet-stream":
            mime = EXT_MIME.get(ext, "application/octet-stream")

        stored = _stored_name(name, ext)
        (upload_dir / stored).write_bytes(data)
        url = f"/api/v1/files/uploads/{stored}"
        kind: Kind = "image" if ext in IMAGE_EXTS else "document"

        item = ProcessedUpload(
            original_name=name,
            stored_name=stored,
            download_url=url,
            kind=kind,
            mime=mime,
        )

        if kind == "image":
            item.extract = f"[Image attached: {name}]"
            if vision_count < MAX_VISION_IMAGES:
                vision = _vision_part(data, mime)
                if vision:
                    item.vision = vision
                    vision_count += 1
                else:
                    item.notes.append("Could not prepare this image for vision.")
            else:
                item.notes.append("Vision input limited to 4 images per turn.")
        else:
            text, err = _extract_document(name, ext, data)
            if err:
                item.error = err
                item.extract = err
            else:
                item.extract = text.strip() or f"(No extractable text in {name})"
            if ext == ".pdf" and not item.error:
                item.figures = _extract_pdf_figures(
                    data, stem=Path(stored).stem, upload_dir=upload_dir
                )
                if item.figures:
                    item.extract = (
                        _format_figure_block(item.figures) + "\n\n" + item.extract
                    )
                    for fig in item.figures[:2]:
                        if vision_count >= MAX_VISION_IMAGES:
                            break
                        part = _vision_part(fig["bytes"], "image/png")
                        if part:
                            item.extra_visions.append(part)
                            vision_count += 1
                    for fig in item.figures:
                        fig.pop("bytes", None)

        processed.append(item)
        logger.info(
            "chat_upload_saved",
            filename=name,
            stored=stored,
            kind=kind,
            bytes=len(data),
        )

    return processed


def build_attachment_query(
    message: str,
    items: list[ProcessedUpload],
) -> tuple[str, list[dict[str, Any]], str]:
    """Return (llm_query, vision_parts, title_source)."""
    user_text = (message or "").strip()
    title_source = user_text or (
        f"📎 {items[0].original_name}" if items else "新会话"
    )
    if not items:
        return user_text, [], title_source

    if not user_text:
        user_text = (
            "请根据附件内容回答 / Please answer based on the attached files."
        )

    dump_parts: list[str] = []
    image_md: list[str] = []
    vision: list[dict[str, Any]] = []
    budget = MAX_EXTRACT_TOTAL

    for item in items:
        header = f"📎 {item.original_name}\n{item.download_url}"
        body = item.error or item.extract
        if item.notes:
            body = (body + "\n" + "\n".join(item.notes)).strip()
        if len(body) > MAX_EXTRACT_PER_FILE:
            body = body[:MAX_EXTRACT_PER_FILE] + "\n…(truncated)"
        if len(body) > budget:
            body = body[:budget] + "\n…(truncated)"
        budget = max(0, budget - len(body))
        dump_parts.append(f"{header}\n\n{body}".rstrip())
        if item.kind == "image":
            image_md.append(f"![{item.original_name}]({item.download_url})")
        for fig in item.figures[:2]:
            url = str(fig.get("url") or "")
            if url:
                image_md.append(f"![{fig.get('label') or 'figure'}]({url})")
        if item.vision:
            vision.append(item.vision)
        vision.extend(item.extra_visions)

    vision = vision[:MAX_VISION_IMAGES]

    query = (
        f"{user_text}\n\n{ATTACH_START}\n"
        + "\n\n".join(dump_parts)
        + f"\n{ATTACH_END}"
    )
    if image_md:
        query += "\n\n" + "\n".join(image_md)

    if vision and not _model_supports_vision():
        query += (
            "\n\n（当前模型不支持识图。图片已保存在对话中；"
            "如需分析图片内容，请在「模型」中切换到支持视觉的模型，"
            "例如 Claude 或 GPT-4o。）"
        )
        vision = []

    if len(query) > MAX_QUERY_CHARS:
        query = query[:MAX_QUERY_CHARS] + "\n…(truncated)"
    return query, vision, title_source


def _model_supports_vision() -> bool:
    try:
        from app.llm.provider_store import resolve_llm
        from app.llm.usage import current_usage_meta

        meta = current_usage_meta()
        cfg = resolve_llm(purpose=meta.purpose, provider_id=meta.provider_id)
    except Exception:
        return False
    name = f"{cfg.kind} {cfg.model or ''}".lower()
    if cfg.kind == "bedrock":
        return any(
            hint in name
            for hint in ("claude", "nova", "sonnet", "haiku", "opus", "gpt-4o")
        )
    return any(hint in name for hint in _VISION_MODEL_HINTS)


def _stored_name(original: str, ext: str) -> str:
    stem = _SAFE_STEM.sub("_", Path(original).stem).strip("._")[:40] or "file"
    return f"{uuid.uuid4().hex[:12]}_{stem}{ext}"


def _format_figure_block(figures: list[dict[str, Any]]) -> str:
    lines = [
        "[PAPER FIGURES — use these src URLs. Do not invent figures with generate_image.]",
        "Resume / 证件照 / 一寸照 / 求职者头像: pick the item labeled portrait and pass it as",
        "hero.image on create_webpage (the cover actually renders that photo).",
        "Paper figures: use layout=figure with image/images.",
    ]
    for i, fig in enumerate(figures, 1):
        lines.append(
            f"{i}. {fig.get('label') or 'figure'} | {fig.get('url')} "
            f"| {fig.get('width')}x{fig.get('height')} | {fig.get('kind')}"
        )
    return "\n".join(lines)


def _pixmap_is_blank(pix) -> bool:
    """Skip solid black/white embeds that are not real photos."""
    samples = getattr(pix, "samples", None) or b""
    if len(samples) < 12:
        return True
    channels = max(1, int(getattr(pix, "n", 3) or 3) - int(getattr(pix, "alpha", 0) or 0))
    step = max(channels, (len(samples) // 2500) * channels)
    vals = samples[0 : len(samples) : step]
    if not vals:
        return True
    return (max(vals) - min(vals)) < 12


def _extract_pdf_figures(
    data: bytes, *, stem: str, upload_dir: Path
) -> list[dict[str, Any]]:
    """Render figure-like PDF pages and keep usable embedded images (ID photos)."""
    try:
        import fitz  # PyMuPDF
    except ImportError:
        logger.info("pymupdf_missing_skip_figures")
        return []

    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        logger.warning("pdf_figure_open_failed", error=str(exc))
        return []

    pages: list[dict[str, Any]] = []
    embeds: list[dict[str, Any]] = []
    try:
        figure_pages: list[int] = []
        for i, page in enumerate(doc):
            text = page.get_text() or ""
            if re.search(r"\b(Figure|Fig\.|Table)\s*\d", text, re.I):
                figure_pages.append(i)
        if not figure_pages:
            figure_pages = list(range(min(3, doc.page_count)))
        elif 0 not in figure_pages:
            figure_pages.insert(0, 0)

        seen_pages: set[int] = set()
        for i in figure_pages:
            if i in seen_pages:
                continue
            seen_pages.add(i)
            page = doc[i]
            pix = page.get_pixmap(matrix=fitz.Matrix(1.7, 1.7), alpha=False)
            if pix.width < 160 or pix.height < 160:
                continue
            blob = pix.tobytes("png")
            stored = f"{stem}_p{i + 1}.png"
            (upload_dir / stored).write_bytes(blob)
            pages.append(
                {
                    "label": f"page {i + 1} render (full page — not an ID photo by itself)",
                    "page": i + 1,
                    "kind": "page",
                    "url": f"/api/v1/files/uploads/{stored}",
                    "width": pix.width,
                    "height": pix.height,
                    "bytes": blob,
                }
            )

        for i, page in enumerate(doc):
            for img in page.get_images(full=True):
                xref = img[0]
                try:
                    pix = fitz.Pixmap(doc, xref)
                    if pix.n - pix.alpha >= 4:
                        pix = fitz.Pixmap(fitz.csRGB, pix)
                    portrait = pix.height >= pix.width * 1.12
                    min_side = 72 if portrait else 240
                    if pix.width < min_side or pix.height < min_side:
                        continue
                    if _pixmap_is_blank(pix):
                        continue
                    blob = pix.tobytes("png")
                except Exception:
                    continue
                stored = f"{stem}_img{xref}.png"
                (upload_dir / stored).write_bytes(blob)
                kind = "portrait" if portrait else "embedded"
                label = (
                    "likely ID/portrait photo — pass as hero.image; do not generate_image"
                    if portrait
                    else f"embedded image on page {i + 1}"
                )
                embeds.append(
                    {
                        "label": label,
                        "page": i + 1,
                        "kind": kind,
                        "url": f"/api/v1/files/uploads/{stored}",
                        "width": pix.width,
                        "height": pix.height,
                        "bytes": blob,
                    }
                )
    finally:
        doc.close()

    portraits = [item for item in embeds if item.get("kind") == "portrait"]
    others = [item for item in embeds if item.get("kind") != "portrait"]
    figures = (portraits + others + pages)[:MAX_PDF_FIGURES]
    logger.info(
        "pdf_figures_extracted",
        count=len(figures),
        portraits=len(portraits),
        stem=stem,
    )
    return figures


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_document(name: str, ext: str, data: bytes) -> tuple[str, str | None]:
    try:
        if ext in {".txt", ".md", ".log", ".xml"}:
            return _decode_text(data), None
        if ext in {".html", ".htm"}:
            return _html_to_text(_decode_text(data)), None
        if ext == ".json":
            return _json_to_text(_decode_text(data)), None
        if ext == ".csv":
            return _csv_to_text(_decode_text(data)), None
        if ext == ".pdf":
            return _pdf_to_text(data), None
        if ext == ".docx":
            return _docx_to_text(data), None
        if ext == ".pptx":
            return _pptx_to_text(data), None
        if ext in {".xlsx", ".xlsm"}:
            return _xlsx_to_text(data), None
        if ext == ".xls":
            return _xls_to_text(data), None
        if ext == ".doc":
            return (
                "",
                f"Cannot read legacy Word .doc ({name}). Please save as .docx and retry.",
            )
        if ext == ".ppt":
            return (
                "",
                f"Cannot read legacy PowerPoint .ppt ({name}). Please save as .pptx and retry.",
            )
    except Exception as exc:
        logger.warning("chat_extract_failed", filename=name, error=str(exc))
        return "", f"Failed to parse '{name}': {exc}"
    return "", f"Unsupported file type for '{name}'."


def _html_to_text(raw: str) -> str:
    text = re.sub(r"(?is)<script[^>]*>.*?</script>", " ", raw)
    text = re.sub(r"(?is)<style[^>]*>.*?</style>", " ", text)
    text = re.sub(r"(?s)<[^>]+>", " ", text)
    text = re.sub(r"&nbsp;", " ", text)
    text = re.sub(r"&amp;", "&", text)
    text = re.sub(r"&lt;", "<", text)
    text = re.sub(r"&gt;", ">", text)
    return re.sub(r"[ \t]+\n", "\n", re.sub(r"\n{3,}", "\n\n", text)).strip()


def _json_to_text(raw: str) -> str:
    try:
        return json.dumps(json.loads(raw), ensure_ascii=False, indent=2)
    except json.JSONDecodeError:
        return raw


def _csv_to_text(raw: str) -> str:
    reader = csv.reader(StringIO(raw))
    rows: list[str] = []
    for i, row in enumerate(reader):
        if i >= 200:
            rows.append("…(truncated after 200 rows)")
            break
        rows.append("\t".join(cell.strip() for cell in row))
    return "\n".join(rows)


def _pdf_to_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    pages: list[str] = []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"--- page {i} ---\n{text}")
    return "\n\n".join(pages) or "(PDF contained no extractable text.)"


def _docx_to_text(data: bytes) -> str:
    from docx import Document

    doc = Document(BytesIO(data))
    chunks: list[str] = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            line = "\t".join(cell.text.strip() for cell in row.cells)
            if line.strip():
                chunks.append(line)
    return "\n".join(chunks)


def _pptx_to_text(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(BytesIO(data))
    slides: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        bits: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text_frame.text.strip()
                if text:
                    bits.append(text)
            if getattr(shape, "has_table", False):
                table = shape.table
                for row in table.rows:
                    line = "\t".join(cell.text.strip() for cell in row.cells)
                    if line.strip():
                        bits.append(line)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            bits.append(f"[Notes]\n{notes}")
        if bits:
            slides.append(f"## Slide {i}\n" + "\n".join(bits))
    return "\n\n".join(slides) or "(Presentation contained no extractable text.)"


def _xlsx_to_text(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    sheets: list[str] = []
    try:
        for sheet_i, ws in enumerate(wb.worksheets):
            if sheet_i >= 5:
                sheets.append("…(further sheets omitted)")
                break
            rows: list[str] = []
            for i, row in enumerate(ws.iter_rows(values_only=True), 1):
                if i > 100:
                    rows.append("…(truncated after 100 rows)")
                    break
                line = "\t".join("" if cell is None else str(cell) for cell in row)
                if line.strip():
                    rows.append(line)
            if rows:
                sheets.append(f"## Sheet {ws.title}\n" + "\n".join(rows))
    finally:
        wb.close()
    return "\n\n".join(sheets) or "(Spreadsheet contained no extractable cells.)"


def _xls_to_text(data: bytes) -> str:
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError(
            "xlrd is not installed; save the workbook as .xlsx and retry."
        ) from exc

    book = xlrd.open_workbook(file_contents=data)
    sheets: list[str] = []
    for sheet_i in range(min(book.nsheets, 5)):
        ws = book.sheet_by_index(sheet_i)
        rows: list[str] = []
        for r in range(min(ws.nrows, 100)):
            line = "\t".join(str(ws.cell_value(r, c)) for c in range(ws.ncols))
            if line.strip():
                rows.append(line)
        if ws.nrows > 100:
            rows.append("…(truncated after 100 rows)")
        if rows:
            sheets.append(f"## Sheet {ws.name}\n" + "\n".join(rows))
    if book.nsheets > 5:
        sheets.append("…(further sheets omitted)")
    return "\n\n".join(sheets) or "(Spreadsheet contained no extractable cells.)"


def _vision_part(data: bytes, mime: str) -> dict[str, Any] | None:
    prepared = _prepare_image(data, mime)
    if prepared is None:
        return None
    blob, out_mime = prepared
    b64 = base64.b64encode(blob).decode("ascii")
    return {
        "type": "image_url",
        "image_url": {"url": f"data:{out_mime};base64,{b64}"},
    }


def _prepare_image(data: bytes, mime: str) -> tuple[bytes, str] | None:
    mime = (mime or "").split(";")[0].strip().lower()
    allowed = {
        "image/jpeg": "image/jpeg",
        "image/jpg": "image/jpeg",
        "image/png": "image/png",
        "image/gif": "image/gif",
        "image/webp": "image/webp",
    }
    if mime in allowed and len(data) <= MAX_VISION_BYTES:
        return data, allowed[mime]
    try:
        from PIL import Image
    except ImportError:
        if mime in allowed:
            return data[:MAX_VISION_BYTES], allowed[mime]
        return None

    try:
        image = Image.open(BytesIO(data))
        image.load()
    except Exception:
        return None

    if getattr(image, "n_frames", 1) > 1:
        image.seek(0)
    if image.mode not in {"RGB", "L"}:
        image = image.convert("RGB")
    elif image.mode == "L":
        image = image.convert("RGB")

    max_side = 2048
    w, h = image.size
    if max(w, h) > max_side:
        scale = max_side / float(max(w, h))
        image = image.resize((max(1, int(w * scale)), max(1, int(h * scale))))

    buf = BytesIO()
    image.save(buf, format="JPEG", quality=85, optimize=True)
    blob = buf.getvalue()
    if len(blob) > MAX_VISION_BYTES:
        buf = BytesIO()
        image.save(buf, format="JPEG", quality=70, optimize=True)
        blob = buf.getvalue()
    if len(blob) > MAX_VISION_BYTES:
        return None
    return blob, "image/jpeg"
